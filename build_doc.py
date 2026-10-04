import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a table cell (values in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    """Set subtle modern borders for a table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:insideV w:val="none"/>'
                            f'<w:left w:val="none"/>'
                            f'<w:right w:val="none"/>'
                            f'</w:tblBorders>')
        tblPr[0].append(borders)

def add_callout(doc, text, title=None, border_color="003366", bg_color="F0F4F8"):
    """Add a shaded callout block with a thick colored left border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=130, bottom=130, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10)
        r = int(border_color[0:2], 16)
        g = int(border_color[2:4], 16)
        b = int(border_color[4:6], 16)
        run_title.font.color.rgb = RGBColor(r, g, b)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.italic = True
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_box(doc, text):
    """Add a monospace diagram or architecture flow box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_custom_table(table, col_widths, headers, data_rows, header_bg="003366"):
    """Format table with corporate styling, zebra striping, and cell margins."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D0D7DE", sz="4")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = Inches(col_widths[c_idx])
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)

def build_troubleshooting_doc():
    doc = docx.Document()
    
    # Page setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "SAP FICO Troubleshooting Handbook: Incidents, Diagnostics & Solutions | Week 5 Deliverable"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Arial"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(128, 128, 128)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Author: Laxmi Gupta  |  SAP Consulting Intern  |  Enterprise Financial Systems Troubleshooting"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.name = "Arial"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    # Document Title Styling
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("SAP FICO Troubleshooting: Common Enterprise Issues, Error Codes, Root Cause Analysis, and Solutions Guide")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(21)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 51, 102) # Corporate Navy

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("An Industrial-Grade Manual for Resolving Financial Accounting & Controlling Incidents, Error Diagnostics, SPRO Customizing Corrections, and S/4HANA Modern Troubleshooting")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(70, 80, 95)

    # Metadata Card Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Author / Candidate:", "Laxmi Gupta (SAP Consulting / ERP Intern)"),
        ("Project Task:", "Week 5 Task: SAP FICO Troubleshooting & Incident Resolution"),
        ("Deliverable Type:", "Formal Industrial Enterprise Troubleshooting Handbook (DOCX)"),
        ("Target System Scope:", "SAP ERP ECC 6.0 EHP8 & SAP S/4HANA 2023 Enterprise Edition"),
        ("Incident Domains Covered:", "FI-GL, FI-AP, FI-AR, FI-AA, FI-BL, CO-CCA, CO-IO, CO-PC, CO-PA, Technical Diagnostics"),
        ("Evaluation Focus:", "Relevance of Issues, Technical Depth, Effectiveness of Solutions, and Clarity")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_val, top=60, bottom=60, left=100, right=100)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGBColor(0, 51, 102)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(40, 40, 40)
        
    set_table_borders(meta_table, color="CBD5E1", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Heading Helpers
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(5)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.bold = True
        r.font.color.rgb = RGBColor(0, 51, 102)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.bold = True
        r.font.color.rgb = RGBColor(30, 64, 120)
        return h

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(9.5)
        r_bold.font.color.rgb = RGBColor(0, 51, 102)
        
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # --- EXECUTIVE SUMMARY ---
    add_h1("Executive Summary")
    add_p(
        "In enterprise resource planning (ERP) environments, high system availability, transactional integrity, and rapid "
        "incident resolution are vital to sustaining business continuity. SAP Financial Accounting (FI) and Controlling (CO) "
        "act as the central accounting and valuation clearinghouse for all enterprise activities. When an error arises—whether "
        "a blocked billing document in Sales & Distribution (SD), an unmapped Goods Receipt in Materials Management (MM), an asset "
        "depreciation failure, or a failed automated payment run—the organization faces immediate operational disruption: month-end "
        "closes freeze, shipments halt, vendor relationships strain, and statutory reporting is delayed."
    )
    add_p(
        "Troubleshooting in SAP FICO is a rigorous, structured engineering discipline. It demands deep understanding of the "
        "delicate interplay between master data maintenance (FS00, BP, CS01) and SPRO customizing rules, cross-module account determination "
        "(OBYC, VKOA, AO90), and technical system logs (ST22, SM21, SM12, SM13)."
    )
    add_callout(
        doc,
        "\"Troubleshooting is root-cause analysis, not patching. Over 80% of SAP FICO production incidents occur at "
        "cross-module integration boundaries (MM-FI, SD-FI, PP-CO). True resolution requires diagnosing whether an issue stems "
        "from master data omissions, customizing table gaps, or process non-compliance—and fixing it at the root.\"",
        title="Core Consulting Axiom",
        border_color="003366",
        bg_color="F0F4F8"
    )

    # --- SECTION 1: 5-STAGE DIAGNOSTIC FRAMEWORK ---
    add_h1("1. The 5-Stage Scientific Diagnostic Framework")
    add_p(
        "To resolve complex FICO incidents without creating secondary regression defects, consultants adhere to a 5-stage framework:"
    )

    add_code_box(
        doc,
        "+-------------------------------------------------------------------------+\n"
        "|                 THE 5-STAGE SAP FICO DIAGNOSTIC FRAMEWORK               |\n"
        "+-------------------+--------------------+--------------------------------+\n"
        "| 1. DECONSTRUCT    | Extract Message Class & Number (SE91); analyze F1 doc. |\n"
        "| 2. CONTEXT TRACE  | Review source vouchers (PO, SO, Billing) & simulate QAS|\n"
        "| 3. CONFLICT AUDIT | Audit Master Data vs. SPRO Customizing tables.         |\n"
        "| 4. SYSTEM LOGS    | Inspect ABAP Dumps (ST22), SysLog (SM21), Locks (SM12).|\n"
        "| 5. RESOLVE & DOC  | Apply fix in DEV, test in QAS, transport, & log in KB. |\n"
        "+-------------------+--------------------+--------------------------------+"
    )

    # --- SECTION 2: FI-GL ISSUES ---
    add_h1("2. General Ledger (FI-GL) & Posting Issues")
    
    table_gl = doc.add_table(rows=6, cols=3)
    t_gl_headers = ["Error Code", "Symptom & Error Text", "Root Cause & Resolution Summary"]
    t_gl_rows = [
        ["F5 286 / 201", "Posting period not open for variant and account type +", "Period locked in OB52. Open Period 1 in OB52 for account type '+' and specific account type."],
        ["KI 235", "Account requires an assignment to a CO object", "Primary cost element missing Cost Center/Order. Enter cost center or configure default in OKB9."],
        ["GLT2 201", "Balancing field 'Profit Center' not filled", "Document splitting issue. Classify GL accounts in GSP_RD; map document types in GSP_VD."],
        ["F5 063", "Exchange rate difference key not defined", "Foreign currency revaluation failure. Maintain realized/unrealized accounts in OBA1 (Key KDF)."],
        ["F5 150 / 151", "Document number range exhausted or not defined", "Number range interval missing for year. Add interval via FBN1 or copy via OBH2."]
    ]
    format_custom_table(table_gl, [1.3, 2.5, 2.7], t_gl_headers, t_gl_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("Detailed Technical Resolutions for General Ledger Incidents:")
    add_bullet(
        "Resolving Posting Period Locks (OB52): ",
        "Posting period variants are assigned to company codes in OBBP. When OB52 is locked, transactions halt. Ensure the global '+' "
        "record and the specific account type (S, K, D, A, M) cover the transaction posting date. If special periods (13-16) are used, "
        "assign an authorization group to restrict year-end auditor adjustments."
    )
    add_bullet(
        "Resolving Missing CO Object Assignments (OKB9 / FS00): ",
        "If an expense GL account is created as Cost Element Category 1, SAP mandates a CO account assignment. For automated SD/MM "
        "postings, maintain transaction OKB9 (Default Account Assignment) by Company Code, Plant, and Profit Center to automatically "
        "derive the cost center without end-user intervention."
    )
    add_bullet(
        "Resolving Document Splitting Errors (GSP_RD / GSP_VD): ",
        "New GL document splitting guarantees a zero-balance balance sheet by Profit Center and Segment. If an account is unclassified, "
        "the posting fails. Use GSP_RD to assign Item Categories (01000 Balance Sheet, 02000 Customer, 03000 Vendor, 20000 Expense) "
        "and ensure zero-balance clearing accounts are defined."
    )

    # --- SECTION 3: FI-AP ISSUES ---
    add_h1("3. Accounts Payable (FI-AP) & Procure-to-Pay (P2P) Issues")

    table_ap = doc.add_table(rows=5, cols=3)
    t_ap_headers = ["Error Code", "Symptom & Error Text", "Root Cause & Resolution Summary"]
    t_ap_rows = [
        ["M8 147", "GR/IR clearing account cannot be determined", "Missing account determination in OBYC for transaction key WRX. Map Valuation Class in OBYC."],
        ["FZ 348 / 349", "Payment run F110: Payment method not found/blocked", "Missing FBZP house bank determination, vendor master missing bank, or invoice blocked with 'A'."],
        ["M8 082 / 084", "Price variance exceeds tolerance limit in MIRO", "Invoice price exceeds PO price beyond OMR6 tolerance keys (PP, BD). Review PO history or release via MRBR."],
        ["7Q 300 / 320", "Withholding tax base amount is zero / code missing", "Vendor master missing WHT recipient flag or WHT type not assigned to Company Code."]
    ]
    format_custom_table(table_ap, [1.3, 2.5, 2.7], t_ap_headers, t_ap_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("Detailed Technical Resolutions for Accounts Payable Incidents:")
    add_bullet(
        "Resolving GR/IR Account Determination (OBYC - WRX): ",
        "During goods receipt (MIGO) and invoice receipt (MIRO), SAP offsets inventory against GR/IR clearing. If transaction OBYC lacks a "
        "mapping for key WRX for the chart of accounts and valuation class, the posting halts. Map a balance sheet GR/IR clearing account in OBYC."
    )
    add_bullet(
        "Resolving F110 Automated Payment Run Failures: ",
        "Check the F110 Proposal Log with 'Additional Log' enabled. If error FZ 349 occurs, inspect transaction FBZP -> Bank Determination -> "
        "Ranking Order and Available Amounts to ensure cash disbursement limits exceed the batch payment amount. Check FBL1N to clear payment blocks."
    )
    add_bullet(
        "Managing Price Variances in MIRO (OMR6 / MRBR): ",
        "When an invoice price deviates from the PO, tolerance key PP in OMR6 blocks payment. Once purchasing confirms the price variance, "
        "the accounting manager executes transaction MRBR (Release Blocked Invoices) to release the item for F110 payment."
    )

    # --- SECTION 4: FI-AR ISSUES ---
    add_h1("4. Accounts Receivable (FI-AR) & Order-to-Cash (O2C) Issues")

    table_ar = doc.add_table(rows=5, cols=3)
    t_ar_headers = ["Error Code", "Symptom & Error Text", "Root Cause & Resolution Summary"]
    t_ar_rows = [
        ["VF 051 / FF 704", "Account determination for SD (VKOA) not possible", "Missing revenue mapping in VKOA for Condition Type (PR00), Sales Org, and Account Assignment Groups."],
        ["V1 150 / Credit", "Customer credit limit exceeded: Order blocked", "Dynamic credit limit check failed in OVA8. Inspect exposure in FD33 / release via VKM3."],
        ["FEBAN Error", "Electronic Bank Statement item not clearing customer", "Search string mismatch or interpretation algorithm in OT83 unable to parse invoice number."],
        ["F-28 Difference", "Difference too large for clearing customer item", "Payment amount does not equal open invoice balance beyond user tolerance limits in OBA3."]
    ]
    format_custom_table(table_ar, [1.3, 2.5, 2.7], t_ar_headers, t_ar_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("Detailed Technical Resolutions for Accounts Receivable Incidents:")
    add_bullet(
        "Resolving SD Revenue Account Determination (VKOA): ",
        "When releasing a billing document (VF02), click Environment -> Account Determination Analysis -> Revenue Accounts. Identify the "
        "unmapped combination of Customer Account Assignment Group (from customer BP) and Material Account Assignment Group (from MM02). "
        "Maintain the missing revenue account in transaction VKOA."
    )
    add_bullet(
        "Resolving Customer Credit Blocks (OVA8 / VKM3): ",
        "Dynamic credit checks aggregate open receivables, deliveries, and orders. If credit limit is breached, credit managers review customer "
        "financials in FD33 / FSCM and release the blocked sales document using transaction VKM3 or VKM1."
    )
    add_bullet(
        "Resolving EBS Bank Clearing Failures (FEBAN / OT83): ",
        "When MT940 statements do not match customer invoices, maintain customized search strings in OTPM/OT83 to extract variable invoice "
        "number formats from the bank statement Note-to-Payee field, or perform interactive manual clearing in FEBAN."
    )

    # --- SECTION 5: FI-AA ISSUES ---
    add_h1("5. Asset Accounting (FI-AA) Issues")

    table_aa = doc.add_table(rows=5, cols=3)
    t_aa_headers = ["Error Code", "Symptom & Error Text", "Root Cause & Resolution Summary"]
    t_aa_rows = [
        ["AA 729 / 688", "Depreciation not completely posted for prior year", "Year-end close AJAB failed due to unposted planned depreciation. Execute AFAB in restart mode."],
        ["AA 400 / 416", "Retirement exceeds NBV / Negative depreciation", "Asset value date is prior to capital acquisition, or depreciation key disallows negative net book value."],
        ["ABST2 Error", "Reconciliation diff between Asset Subledger and GL", "Manual journal posted to reconciliation account. Reconcile with RAABST02 and set Recon flag in FS00."],
        ["AIBU Error", "Asset under Construction (AuC) settlement failure", "AuC capital project lacks settlement rule. Maintain investment profile or settlement rule in AIAB."]
    ]
    format_custom_table(table_aa, [1.3, 2.5, 2.7], t_aa_headers, t_aa_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("Detailed Technical Resolutions for Asset Accounting Incidents:")
    add_bullet(
        "Resolving Year-End Asset Closing Errors (AJAB / AFAB): ",
        "SAP prevents closing an asset fiscal year if planned depreciation remains unposted. Execute AFBP to identify unposted assets, "
        "run AFAB in background mode for Period 12 in 'Repeat' or 'Restart' mode, and re-execute AJAB to close the asset year."
    )
    add_bullet(
        "Resolving Subledger vs. General Ledger Differences (ABST2): ",
        "If report ABST2 flags discrepancies, audit the account in FBL3N. Ensure the GL account field 'Recon. account for acct type' is "
        "set to 'Assets' in FS00 to prohibit manual FB50 postings, and post adjusting entries via ABF1."
    )

    # --- SECTION 6: CO MANAGERIAL ISSUES ---
    add_h1("6. Controlling (CO) Managerial & Costing Issues")

    table_co = doc.add_table(rows=5, cols=3)
    t_co_headers = ["Error Code", "Symptom & Error Text", "Root Cause & Resolution Summary"]
    t_co_rows = [
        ["KD 205", "Order requires settlement rule or rule not maintained", "Internal or production order settlement KO88/CO88 failed. Maintain settlement rule in KO02/CO02."],
        ["CK 465 / KI 113", "No standard cost estimate found for material", "Material has price control 'S' but lacks released cost estimate. Calculate in CK11N; mark/release in CK24."],
        ["C+ 704", "Material Ledger: Closing entries cannot be posted", "Period shift MMPV out of sync with CKMLCP. Follow strict sequential closing steps in CKMLCP."],
        ["KE 015 / 083", "Profitability Analysis characteristic derivation failed", "Derivation rule missing in KEDR. Add fallback rule or table lookup for product/customer characteristics."]
    ]
    format_custom_table(table_co, [1.3, 2.5, 2.7], t_co_headers, t_co_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("Detailed Technical Resolutions for Controlling Incidents:")
    add_bullet(
        "Maintaining Order Settlement Rules (KO02 / CO02): ",
        "Orders accumulating costs require a settlement rule to transfer balances at month-end. In KO02/CO02, navigate to Settlement Rule, "
        "specify category (CTR Cost Center, PSG Profitability Segment), receiver number, and 100% distribution."
    )
    add_bullet(
        "Marking and Releasing Standard Costs (CK11N / CK24): ",
        "When standard price 'S' materials lack cost estimates, sales deliveries and goods issues fail. Calculate costs via CK11N (or CK40N "
        "for mass runs), activate marking allowance in CK24, mark the future price, and release to current standard price."
    )

    # --- SECTION 7: TECHNICAL DIAGNOSTICS ---
    add_h1("7. Enterprise Technical Diagnostics & System Tools")
    add_p(
        "When application errors do not provide conclusive root causes, functional consultants deploy standard system diagnostic utilities:"
    )

    table_tools = doc.add_table(rows=7, cols=3)
    t_tools_headers = ["T-Code", "Diagnostic Tool Name", "Technical Usage & Investigation Focus"]
    t_tools_rows = [
        ["ST22", "ABAP Dump Analysis", "Investigates fatal runtime crashes, exceptions, and exact failing source code line numbers."],
        ["SM21", "System Log", "Audits database connection drops, memory allocation faults, and system-wide error events."],
        ["SM12", "Lock Entries Inspection", "Identifies exclusive table locks blocking concurrent users; allows releasing orphaned locks."],
        ["SM13", "Update Task Monitor", "Diagnoses failed asynchronous database update tasks (VBHDR) and re-triggers updates."],
        ["SE16N", "Table Browser", "Direct inspection of relational database records in BKPF, BSEG, ACDOCA, and master tables."],
        ["ST05", "Performance SQL Trace", "Analyzes query execution times, buffer hits, and identifies missing database indexes."]
    ]
    format_custom_table(table_tools, [1.0, 2.2, 3.3], t_tools_headers, t_tools_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- SECTION 8: ROOT CAUSE PREVENTION ---
    add_h1("8. Root Cause Prevention & Governance Best Practices")
    add_p(
        "To establish high enterprise stability and minimize recurring helpdesk tickets:"
    )
    add_bullet("1. Automated Cutover Calendar Governance: ", "Automate the opening of posting periods (OB52) and material periods (MMPV) using scheduled background jobs to prevent first-day-of-the-month transaction rejections.")
    add_bullet("2. Master Data Governance (SAP MDG): ", "Implement mandatory validation gates in master data creation workflows. Prohibit releasing G/L accounts without document splitting classification, and mandate bank/tax validation on vendor master records.")
    add_bullet("3. Strict Transport Freeze & Quality Validation: ", "Enforce zero configuration directly in Production. All customizing must originate in Development (DEV), undergo unit and integration testing in Quality (QAS), and be transported via documented Transport Requests.")
    add_bullet("4. Proactive System Health Monitoring: ", "Set up automated daily alerts for failed update tasks (SM13), ABAP short dumps (ST22), and orphaned lock entries (SM12) to catch anomalies before end-users report them.")

    # Save document
    output_filename = "SAP_FICO_Troubleshooting_Guide.docx"
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), output_filename)
    doc.save(output_path)
    print(f"Successfully generated SAP FICO Troubleshooting Guide at: {output_path}")

if __name__ == "__main__":
    build_troubleshooting_doc()
